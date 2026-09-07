from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.lead_service import LeadService


HEADER_FILL = PatternFill("solid", fgColor="11191E")
HEADER_FONT = Font(color="EDF2F4", bold=True)
ACCENT_FILL = PatternFill("solid", fgColor="C9F24A")


def _style_table(sheet: object) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24
    for column in sheet.columns:
        values = [str(cell.value or "") for cell in column]
        width = min(max(len(value) for value in values) + 3, 42)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = max(width, 12)


async def export_workbook(service: LeadService, output: str | Path) -> Path:
    destination = Path(output)
    destination.parent.mkdir(parents=True, exist_ok=True)
    leads = await service.list_all()
    history = await service.list_history()
    stats = await service.get_stats()

    workbook = Workbook()
    leads_sheet = workbook.active
    leads_sheet.title = "Заявки"
    leads_sheet.append(
        [
            "ID",
            "Создана",
            "Клиент",
            "Услуга",
            "Телефон",
            "Удобное время",
            "Статус",
            "Менеджер",
            "Источник",
            "Комментарий",
        ]
    )
    for lead in leads:
        leads_sheet.append(
            [
                lead.id,
                lead.created_at.isoformat(),
                lead.name,
                lead.service_type,
                lead.phone,
                lead.preferred_time,
                lead.status.value,
                lead.manager_id,
                lead.source,
                lead.comment,
            ]
        )
    _style_table(leads_sheet)

    history_sheet = workbook.create_sheet("История")
    history_sheet.append(["ID", "Заявка", "Изменено", "Из", "В", "Исполнитель"])
    for item in history:
        history_sheet.append(
            [
                item.id,
                item.lead_id,
                item.changed_at.isoformat(),
                item.from_status.value if item.from_status else "",
                item.to_status.value,
                item.actor_id,
            ]
        )
    _style_table(history_sheet)

    summary_sheet = workbook.create_sheet("Сводка")
    summary_sheet.append(["Показатель", "Значение"])
    summary_sheet.append(["Всего заявок", stats.total])
    summary_sheet.append(["Новые", stats.by_status["new"]])
    summary_sheet.append(["В работе", stats.by_status["in_progress"]])
    summary_sheet.append(["Отложены", stats.by_status["postponed"]])
    summary_sheet.append(["Закрыты", stats.by_status["done"]])
    summary_sheet.append(["Просрочено новых", stats.overdue_new])
    _style_table(summary_sheet)
    summary_sheet["B2"].fill = ACCENT_FILL
    summary_sheet["B2"].font = Font(color="080B0D", bold=True)

    workbook.save(destination)
    return destination

