from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app.models import LeadCreate, LeadStatus
from app.services.export_service import export_workbook
from app.services.lead_service import LeadService


@pytest.mark.asyncio
async def test_excel_export_contains_leads_history_and_summary(
    service: LeadService, tmp_path: Path
) -> None:
    lead = await service.create_lead(
        LeadCreate(
            name="Анна Крылова",
            phone="+7 999 123-45-67",
            service_type="CRM",
            comment="Нужна очередь менеджеров",
            preferred_time="Утро",
            source="telegram",
        )
    )
    await service.change_status(lead.id, LeadStatus.IN_PROGRESS, actor_id=7001)
    output = tmp_path / "leads.xlsx"

    result = await export_workbook(service, output)
    workbook = openpyxl.load_workbook(result, data_only=False)

    assert workbook.sheetnames == ["Заявки", "История", "Сводка"]
    assert workbook["Заявки"]["A2"].value == lead.id
    assert workbook["Заявки"]["D2"].value == "CRM"
    assert workbook["Заявки"]["G2"].value == "in_progress"
    assert workbook["История"].max_row == 3
    assert workbook["Сводка"]["B2"].value == 1

